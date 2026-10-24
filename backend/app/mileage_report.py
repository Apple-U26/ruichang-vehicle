"""Read-only mileage snapshots and self-contained Excel reports."""

import io
import warnings
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Literal
from urllib.parse import unquote

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field, model_validator
from sqlalchemy import or_, select

from app.access import apply_vehicle_scope, ensure_vehicle_access
from app.models import MileageRecord, Vehicle

MAX_ROWS = 5000
MAX_PHOTOS = 400
MAX_IMAGE_BYTES = 10 * 1024 * 1024
MAX_EMBED_BYTES = 32 * 1024 * 1024
HEADER_ROW = 6


class ReportRequest(BaseModel):
    date_from: date
    date_to: date
    time_basis: Literal['departure', 'return'] = 'departure'
    vehicle_id: int | None = Field(default=None, gt=0)
    driver_name: str = Field(default='', max_length=50)
    keyword: str = Field(default='', max_length=100)
    status: Literal['OUT', 'CLOSED'] | None = None
    photos: Literal['none', 'representative', 'all'] = 'representative'

    @model_validator(mode='after')
    def valid_period(self):
        if self.date_to < self.date_from:
            raise ValueError('结束日期不能早于开始日期')
        if (self.date_to - self.date_from).days >= 366:
            raise ValueError('单次导出最多366天，请分段导出')
        self.driver_name = self.driver_name.strip()
        return self


def apply_period(stmt, date_from=None, date_to=None, time_basis='departure'):
    if time_basis not in {'departure', 'return'}:
        raise HTTPException(status_code=400, detail='时间依据不正确')
    if date_from and date_to and date_from > date_to:
        raise HTTPException(status_code=400, detail='结束日期不能早于开始日期')
    column = MileageRecord.close_time if time_basis == 'return' else MileageRecord.trip_date
    if date_from:
        stmt = stmt.where(column >= datetime.combine(date_from, time.min))
    if date_to:
        if date_to == date.max:
            raise HTTPException(status_code=400, detail='结束日期超出范围')
        stmt = stmt.where(column < datetime.combine(date_to + timedelta(days=1), time.min))
    return stmt


def photo_urls(value):
    return [part.strip() for part in (value or '').split(',') if part.strip()]


def selected_photos(row, mode):
    if mode == 'none':
        return []
    result = []
    for key, label in [('out_photo', '出车'), ('in_photo', '收车')]:
        urls = photo_urls(row[key])
        chosen = urls[:1] if mode == 'representative' else urls
        result.extend((label, i + 1, len(urls), url) for i, url in enumerate(chosen))
    return result


def collect_rows(db, user, filters: ReportRequest):
    if filters.vehicle_id:
        ensure_vehicle_access(user, db.get(Vehicle, filters.vehicle_id))
    stmt = select(MileageRecord, Vehicle.plate_no).join(Vehicle, Vehicle.id == MileageRecord.vehicle_id)
    stmt = apply_vehicle_scope(user, stmt)
    stmt = apply_period(stmt, filters.date_from, filters.date_to, filters.time_basis)
    if filters.vehicle_id:
        stmt = stmt.where(MileageRecord.vehicle_id == filters.vehicle_id)
    if filters.status:
        stmt = stmt.where(MileageRecord.status == filters.status)
    if filters.keyword.strip():
        like = '%' + filters.keyword.strip() + '%'
        stmt = stmt.where(or_(Vehicle.plate_no.like(like), MileageRecord.driver_name.like(like), MileageRecord.purpose.like(like)))
    records = db.execute(stmt.order_by(MileageRecord.trip_date, MileageRecord.id).limit(MAX_ROWS + 1)).all()
    if len(records) > MAX_ROWS:
        raise HTTPException(status_code=400, detail='记录超过5000条，请缩小日期或车辆范围')
    rows = []
    for record, plate in records:
        name = (record.driver_name or '').strip()
        if filters.driver_name and name != filters.driver_name:
            continue
        row = {key: getattr(record, key) for key in [
            'id', 'vehicle_id', 'trip_date', 'close_time', 'out_mileage', 'in_mileage',
            'distance', 'departure', 'destination', 'purpose', 'out_photo', 'in_photo',
            'status', 'abnormal', 'abnormal_reason', 'remark',
        ]}
        row.update(plate_no=plate, driver_name=name)
        rows.append(row)
    photo_count = sum(len(selected_photos(row, filters.photos)) for row in rows)
    if photo_count > MAX_PHOTOS:
        raise HTTPException(status_code=400, detail='所选照片超过400张，请选择代表照片、无照片或缩小日期范围')
    return rows


def classify(row):
    if row['status'] != 'CLOSED':
        return '未收车', Decimal('0'), Decimal('0')
    values = [row['out_mileage'], row['in_mileage'], row['distance']]
    if any(value is None or not Decimal(str(value)).is_finite() for value in values):
        return '数据待核实', Decimal('0'), Decimal('0')
    out, returned, distance = map(lambda value: Decimal(str(value)), values)
    if min(out, returned, distance) < 0 or returned < out or abs(returned - out - distance) > Decimal('0.1'):
        return '数据待核实', Decimal('0'), Decimal('0')
    if row['abnormal']:
        return '异常待核实', Decimal('0'), distance
    return '有效里程', distance, Decimal('0')


def summarize(rows):
    groups = {}
    for row in rows:
        key = row['driver_name'] or None
        group = groups.setdefault(key, {'name': key or '未填写姓名', 'trips': 0, 'closed': 0,
            'valid': Decimal('0'), 'abnormal': Decimal('0'), 'open': 0, 'invalid': 0,
            'vehicles': set(), 'days': set()})
        category, valid, abnormal = classify(row)
        group['trips'] += 1
        group['closed'] += row['status'] == 'CLOSED'
        group['open'] += row['status'] != 'CLOSED'
        group['invalid'] += category == '数据待核实'
        group['valid'] += valid
        group['abnormal'] += abnormal
        group['vehicles'].add(row['vehicle_id'])
        group['days'].add(row['trip_date'].date())
    return sorted(groups.values(), key=lambda group: (group['name'] == '未填写姓名', group['name']))


def literal(cell, value):
    cell.value = value
    # Exported user text must remain text, including strings beginning with '='.
    if isinstance(value, str):
        cell.data_type = 's'
    if isinstance(value, datetime):
        cell.number_format = 'yyyy-mm-dd hh:mm'
    elif isinstance(value, (int, float, Decimal)):
        cell.number_format = '0.0' if isinstance(value, (float, Decimal)) else '0'
    cell.alignment = Alignment(vertical='center', wrap_text=True)


def setup_sheet(ws, title, params, headers, widths):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(1, 1, title).font = Font(name='Microsoft YaHei', size=17, bold=True, color='183B56')
    basis = '出车时间' if params['time_basis'] == 'departure' else '收车时间'
    notes = [f"日期：{params['date_from']} 至 {params['date_to']}（含结束日全天）｜统计依据：{basis}",
        '同名合并，姓名前后空格忽略；异常、数据不完整及未收车行程不计有效里程。',
        '时间按系统既有登记值展示；收车时间为登记时间，历史导入可能只有月份，缺失时间留空。',
        f"驾驶人：{params.get('driver_name') or '全部'}｜照片：{ {'none':'不附图','representative':'出收车各一张代表图','all':'全部照片'}[params['photos']] }｜生成：{datetime.now():%Y-%m-%d %H:%M}"]
    for index, note in enumerate(notes, 2):
        ws.merge_cells(start_row=index, start_column=1, end_row=index, end_column=len(headers))
        literal(ws.cell(index, 1), note)
        ws.cell(index, 1).font = Font(name='Microsoft YaHei', size=10, color='52667A')
        ws.row_dimensions[index].height = 22
    for col, heading in enumerate(headers, 1):
        cell = ws.cell(HEADER_ROW, col, heading)
        cell.font = Font(name='Microsoft YaHei', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='245A81')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    ws.row_dimensions[HEADER_ROW].height = 32
    ws.freeze_panes = f'D{HEADER_ROW + 1}'
    ws.print_title_rows = f'1:{HEADER_ROW}'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def thumbnail(url, upload_path):
    from PIL import Image, ImageOps
    from openpyxl.drawing.image import Image as ExcelImage
    if not url.startswith('/uploads/'):
        return None, '附件路径不受支持'
    name = unquote(url[len('/uploads/'):])
    if not name or Path(name).name != name or '/' in name or '\\' in name or '\x00' in name:
        return None, '附件路径无效'
    root = Path(upload_path).resolve()
    path = (root / name).resolve()
    if not path.is_relative_to(root) or not path.is_file():
        return None, '文件缺失'
    if path.stat().st_size > MAX_IMAGE_BYTES:
        return None, '原图过大，请在系统中查看'
    try:
        with warnings.catch_warnings():
            warnings.simplefilter('error', Image.DecompressionBombWarning)
            with Image.open(path) as original:
                if original.width * original.height > 25_000_000:
                    return None, '图片像素过大，请在系统中查看'
                # Decode at reduced resolution where supported, preserving originals.
                original.draft('RGB', (1200, 900))
                image = ImageOps.exif_transpose(original)
                image.thumbnail((960, 720))
                if image.mode != 'RGB':
                    background = Image.new('RGB', image.size, 'white')
                    background.paste(image, mask=image.getchannel('A') if 'A' in image.getbands() else None)
                    image = background
                buffer = io.BytesIO()
                image.save(buffer, format='JPEG', quality=78, optimize=True)
                buffer.seek(0)
                picture = ExcelImage(buffer)
                ratio = min(320 / picture.width, 200 / picture.height)
                picture.width, picture.height = picture.width * ratio, picture.height * ratio
                return (picture, buffer), '已嵌入压缩图片'
    except (OSError, ValueError, Image.DecompressionBombError, Image.DecompressionBombWarning):
        return None, '无法读取的图片或非图片附件'


def build_workbook(rows, params, upload_path, output_path, progress=lambda value: None):
    wb = Workbook()
    summary, detail, photos = wb.active, wb.create_sheet('里程明细'), wb.create_sheet('照片凭证')
    summary.title = '个人里程汇总'
    setup_sheet(summary, '个人里程汇总', params,
        ['姓名', '出车次数', '已收车次数', '有效里程(km)', '异常待核实里程(km)', '未收车次数', '涉及车辆数', '出车天数', '数据待核实次数'],
        [20, 14, 16, 19, 25, 16, 16, 14, 20])
    groups = summarize(rows)
    for index, group in enumerate(groups, HEADER_ROW + 1):
        values = [group['name'], group['trips'], group['closed'], float(group['valid']), float(group['abnormal']),
            group['open'], len(group['vehicles']), len(group['days']), group['invalid']]
        for col, value in enumerate(values, 1):
            literal(summary.cell(index, col), value)
    total_row = HEADER_ROW + len(groups) + 1
    totals = ['合计', len(rows), sum(r['status'] == 'CLOSED' for r in rows),
        float(sum((g['valid'] for g in groups), Decimal('0'))), float(sum((g['abnormal'] for g in groups), Decimal('0'))),
        sum(g['open'] for g in groups), len({r['vehicle_id'] for r in rows}),
        len({r['trip_date'].date() for r in rows}), sum(g['invalid'] for g in groups)]
    for col, value in enumerate(totals, 1):
        literal(summary.cell(total_row, col), value)
        summary.cell(total_row, col).font = Font(bold=True, color='183B56')
        summary.cell(total_row, col).fill = PatternFill('solid', fgColor='DCEAF5')
    summary.auto_filter.ref = f'A{HEADER_ROW}:I{max(HEADER_ROW, total_row - 1)}'
    setup_sheet(detail, '里程记录明细', params,
        ['行程编号', '车牌号', '驾驶人', '出车时间', '收车时间', '出车里程(km)', '收车里程(km)', '本次里程(km)',
         '有效里程(km)', '异常待核实里程(km)', '统计归类', '出车照片', '收车照片', '出发地', '目的地', '事由', '异常说明', '备注'],
        [16, 16, 18, 23, 23, 19, 19, 19, 19, 25, 19, 23, 23, 22, 22, 32, 32, 32])
    setup_sheet(photos, '出收车照片凭证', params,
        ['行程编号', '车牌号', '驾驶人', '照片类型', '张数', '对应登记时间', '照片', '说明'],
        [16, 16, 18, 14, 14, 23, 48, 34])
    photo_row, embedded_bytes, buffers, missing = HEADER_ROW + 1, 0, [], 0
    for index, row in enumerate(rows, HEADER_ROW + 1):
        category, valid, abnormal = classify(row)
        values = [f"LC-{row['id']:06d}", row['plate_no'], row['driver_name'] or '未填写姓名', row['trip_date'], row['close_time'],
            float(row['out_mileage']) if row['out_mileage'] is not None else None,
            float(row['in_mileage']) if row['in_mileage'] is not None else None,
            float(row['distance']) if row['distance'] is not None else None, float(valid), float(abnormal), category,
            '', '', row['departure'] or '', row['destination'] or '', row['purpose'] or '', row['abnormal_reason'] or '', row['remark'] or '']
        for col, value in enumerate(values, 1):
            literal(detail.cell(index, col), value)
        detail.row_dimensions[index].height = 30
        targets = {}
        for label, position, count, url in selected_photos(row, params['photos']):
            targets.setdefault(label, photo_row)
            image, note = thumbnail(url, upload_path)
            data = [values[0], row['plate_no'], values[2], label, f'{position}/{count}',
                row['trip_date'] if label == '出车' else row['close_time'], '', note]
            for col, value in enumerate(data, 1):
                literal(photos.cell(photo_row, col), value)
            if image:
                picture, buffer = image
                embedded_bytes += buffer.getbuffer().nbytes
                if embedded_bytes > MAX_EMBED_BYTES:
                    raise ValueError('照片压缩后仍超过32MB，请缩小范围或选择代表照片')
                buffers.append(buffer)
                photos.add_image(picture, f'G{photo_row}')
                photos.row_dimensions[photo_row].height = 155
            else:
                missing += 1
                photos.row_dimensions[photo_row].height = 34
            photos.cell(photo_row, 1).hyperlink = f"#'里程明细'!A{index}"
            photo_row += 1
        for col, key, label in [(12, 'out_photo', '出车'), (13, 'in_photo', '收车')]:
            count = len(photo_urls(row[key]))
            cell = detail.cell(index, col)
            literal(cell, f'{count}张 · 查看凭证' if label in targets else f'{count}张（未附图）' if count else '未上传')
            if label in targets:
                cell.hyperlink = f"#'照片凭证'!A{targets[label]}"
                cell.font = Font(color='1769AA', underline='single')
        progress(int((index - HEADER_ROW) / max(len(rows), 1) * 85))
    if not rows:
        literal(detail.cell(HEADER_ROW + 1, 1), '所选条件下暂无记录')
    if photo_row == HEADER_ROW + 1:
        literal(photos.cell(photo_row, 1), '本次未附照片，或所选记录未上传照片')
    detail.auto_filter.ref = f'A{HEADER_ROW}:R{max(HEADER_ROW, HEADER_ROW + len(rows))}'
    wb.active = 0
    try:
        wb.save(output_path)
    finally:
        for buffer in buffers:
            buffer.close()
        wb.close()
    progress(100)
    return {'records': len(rows), 'people': len(groups), 'photos': photo_row - HEADER_ROW - 1, 'unreadable_photos': missing,
        'valid_mileage': float(sum((g['valid'] for g in groups), Decimal('0'))),
        'abnormal_mileage': float(sum((g['abnormal'] for g in groups), Decimal('0')))}
