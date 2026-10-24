import hashlib
import io
import json
import os
import tempfile
import time
import unittest
import zipfile
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import test_safety as baseline
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import select

from app import report_jobs
from app.database import SessionLocal
from app.mileage_report import HEADER_ROW, ReportRequest, build_workbook, collect_rows, summarize, thumbnail
from app.models import MileageRecord, User


class MileageReportTests(unittest.TestCase):
    request = baseline.SafetyTests.request

    def setUp(self):
        self.scratch = tempfile.TemporaryDirectory()
        os.environ['MILEAGE_REPORT_DIR'] = str(Path(self.scratch.name) / 'jobs')
        baseline.SafetyTests.setUp(self)
        self.uploads = Path(baseline.UPLOADS.name)
        self.image_names = ['report-test-out.jpg', 'report-test-second.jpg', 'report-test-in.png']
        for index, name in enumerate(self.image_names):
            Image.new('RGB', (1200, 800), (60 + index * 40, 100, 160)).save(self.uploads / name)
        self.image_hashes = {name: hashlib.sha256((self.uploads / name).read_bytes()).hexdigest() for name in self.image_names}
        self.record_ids = []
        with SessionLocal() as db:
            specs = [
                ('A', '2026-10-01T08:00:00', '2026-10-01T09:00:00', 100, 110, 10, '张三', False),
                ('B', '2026-10-02T08:00:00', '2026-10-02T09:00:00', 5, 25, 20, ' 张三 ', False),
                ('C', '2026-10-31T23:59:59', '2026-11-01T01:00:00', 0, 15, 15, '李四', False),
                ('D', '2026-11-01T00:00:00', '2026-11-01T01:00:00', 0, 9, 9, '十一月', False),
                ('A', '2026-10-03T08:00:00', None, 110, None, 0, '李四', False),
                ('A', '2026-10-04T08:00:00', '2026-10-04T09:00:00', 120, 150, 30, '张三', True),
                ('A', '2026-10-05T08:00:00', '2026-10-05T09:00:00', 150, 140, -10, '张三', True),
                ('A', '2026-10-06T08:00:00', '2026-10-06T09:00:00', 200, 205, 5, None, False),
                ('A', '2026-10-07T08:00:00', '2026-10-07T09:00:00', 300, 305, 5, '=SUM(1,2)', False),
            ]
            for index, (vehicle, start, end, out, returned, distance, name, abnormal) in enumerate(specs):
                row = MileageRecord(vehicle_id=self.vehicles[vehicle], trip_date=datetime.fromisoformat(start),
                    close_time=datetime.fromisoformat(end) if end else None, out_mileage=out, in_mileage=returned,
                    distance=distance, driver_name=name, status='CLOSED' if end else 'OUT', abnormal=abnormal,
                    out_photo='/uploads/report-test-out.jpg,/uploads/report-test-second.jpg' if index == 0 else '/uploads/missing.jpg' if index == 1 else None,
                    in_photo='/uploads/report-test-in.png' if index == 0 else None, created_by=self.ids['admin'])
                db.add(row)
                db.flush()
                self.record_ids.append(row.id)
            db.commit()

    def tearDown(self):
        baseline.SafetyTests.tearDown(self)
        os.environ.pop('MILEAGE_REPORT_DIR', None)
        self.scratch.cleanup()

    def filters(self, **kwargs):
        return ReportRequest(date_from='2026-10-01', date_to='2026-10-31', **kwargs)

    def snapshot(self, role='admin', **kwargs):
        filters = self.filters(**kwargs)
        with SessionLocal() as db:
            rows = collect_rows(db, db.get(User, self.ids[role]), filters)
        return rows, filters

    def build(self, **kwargs):
        rows, filters = self.snapshot(**kwargs)
        path = Path(self.scratch.name) / 'report.xlsx'
        stats = build_workbook(rows, filters.model_dump(mode='json'), self.uploads, path)
        return path, stats

    def test_period_includes_end_day_and_handles_cross_month_basis(self):
        october = self.request('GET', '/mileages?date_from=2026-10-01&date_to=2026-10-31').json()
        self.assertIn(self.record_ids[2], [row['id'] for row in october])
        self.assertNotIn(self.record_ids[3], [row['id'] for row in october])
        returned = self.request('GET', '/mileages?month=2026-11&time_basis=return').json()
        departed = self.request('GET', '/mileages?month=2026-11').json()
        self.assertEqual({r['id'] for r in returned}, {self.record_ids[2], self.record_ids[3]})
        self.assertEqual({r['id'] for r in departed}, {self.record_ids[3]})

    def test_same_name_across_vehicles_and_missing_name_bucket(self):
        rows, _ = self.snapshot()
        groups = {g['name']: g for g in summarize(rows)}
        self.assertEqual(len(rows), 8)
        self.assertEqual(groups['张三']['valid'], 30)
        self.assertEqual(groups['张三']['abnormal'], 30)
        self.assertEqual(groups['张三']['invalid'], 1)
        self.assertEqual(len(groups['张三']['vehicles']), 2)
        self.assertEqual(groups['李四']['open'], 1)
        self.assertEqual(groups['未填写姓名']['valid'], 5)

    def test_exact_name_filter_strips_outer_spaces_only(self):
        rows, _ = self.snapshot(driver_name=' 张三 ')
        self.assertEqual(len(rows), 4)
        self.assertTrue(all(row['driver_name'] == '张三' for row in rows))
        self.assertEqual(len(self.snapshot(driver_name='张 三')[0]), 0)

    def test_scope_applies_to_detail_summary_and_options(self):
        rows, _ = self.snapshot(role='pm_a')
        self.assertTrue(all(row['vehicle_id'] == self.vehicles['A'] for row in rows))
        self.assertEqual(sum(g['valid'] for g in summarize(rows)), 20)
        options = self.request('GET', '/mileage-reports/options', 'pm_a').json()
        self.assertEqual([row['id'] for row in options['vehicles']], [self.vehicles['A']])
        response = self.request('POST', '/mileage-reports', 'pm_a', json={**self.filters(photos='none').model_dump(mode='json'), 'vehicle_id': self.vehicles['B']})
        self.assertEqual(response.status_code, 403, response.text)

    def test_workbook_native_datetimes_reconciliation_and_text_safety(self):
        path, stats = self.build(photos='none')
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ['个人里程汇总', '里程明细', '照片凭证'])
        details, summary = wb['里程明细'], wb['个人里程汇总']
        self.assertIsInstance(details.cell(HEADER_ROW + 1, 4).value, datetime)
        self.assertIsInstance(details.cell(HEADER_ROW + 1, 5).value, datetime)
        self.assertEqual(details.cell(HEADER_ROW + 1, 4).number_format, 'yyyy-mm-dd hh:mm')
        self.assertEqual(sum(details.cell(r, 9).value for r in range(HEADER_ROW + 1, details.max_row + 1)), 55)
        self.assertEqual(summary.cell(summary.max_row, 4).value, 55)
        self.assertEqual(stats['valid_mileage'], 55)
        formula_cells = [c for row in details for c in row if c.value == '=SUM(1,2)']
        self.assertTrue(formula_cells)
        self.assertTrue(all(c.data_type == 's' for c in formula_cells))
        self.assertTrue(all(c.hyperlink is None or (c.hyperlink.target is None and c.hyperlink.location) for row in details for c in row))
        wb.close()

    def test_representative_and_all_photos_are_embedded_without_altering_originals(self):
        for mode, expected in [('representative', 2), ('all', 3), ('none', 0)]:
            path, stats = self.build(photos=mode)
            with zipfile.ZipFile(path) as archive:
                self.assertEqual(len([name for name in archive.namelist() if name.startswith('xl/media/')]), expected)
            self.assertEqual(stats['unreadable_photos'], 1 if mode != 'none' else 0)
        for name, digest in self.image_hashes.items():
            self.assertEqual(hashlib.sha256((self.uploads / name).read_bytes()).hexdigest(), digest)

    def test_photo_path_cannot_escape_upload_root_or_fetch_remote_urls(self):
        for url in ['/uploads/../secret.jpg', '/uploads/%2e%2e/secret.jpg', 'https://example.com/photo.jpg', '/uploads/..\\secret.jpg']:
            picture, message = thumbnail(url, self.uploads)
            self.assertIsNone(picture, url)
            self.assertTrue(message)

    def test_photo_limit_is_checked_before_scheduling(self):
        with patch('app.mileage_report.MAX_PHOTOS', 1):
            response = self.request('POST', '/mileage-reports', json=self.filters().model_dump(mode='json'))
            self.assertEqual(response.status_code, 400, response.text)

    def test_date_validation_and_driver_permission(self):
        for changes, expected in [({'date_to': '2026-09-01'}, 422), ({'date_to': '2028-01-01'}, 422), ({'time_basis': 'wrong'}, 422)]:
            response = self.request('POST', '/mileage-reports', json={**self.filters().model_dump(mode='json'), **changes})
            self.assertEqual(response.status_code, expected, response.text)
        self.assertEqual(self.request('GET', '/mileages?date_from=2026-10-24&date_to=2026-10-01').status_code, 400)
        for path in ['/mileage-reports/options', '/mileage-reports/' + 'a' * 32]:
            self.assertEqual(self.request('GET', path, 'driver').status_code, 403)
        self.assertEqual(self.request('POST', '/mileage-reports', 'driver', json=self.filters().model_dump(mode='json')).status_code, 403)

    def test_background_generation_download_owner_and_current_scope(self):
        response = self.request('POST', '/mileage-reports', json=self.filters().model_dump(mode='json'))
        self.assertEqual(response.status_code, 200, response.text)
        job_id = response.json()['id']
        self.assertEqual(self.request('GET', '/mileage-reports/' + job_id, 'finance').status_code, 404)
        for _ in range(300):
            self.assertEqual(self.request('GET', '/health').status_code, 200)
            status = self.request('GET', '/mileage-reports/' + job_id).json()
            if status['status'] in ['ready', 'failed']:
                break
            time.sleep(0.1)
        self.assertEqual(status['status'], 'ready', status)
        response = self.request('GET', f'/mileage-reports/{job_id}/download')
        self.assertEqual(response.status_code, 200, response.text[:300] if response.status_code != 200 else '')
        self.assertEqual(response.headers['cache-control'], 'private, no-store')
        self.assertEqual(load_workbook(io.BytesIO(response.content)).sheetnames, ['个人里程汇总', '里程明细', '照片凭证'])
        with SessionLocal() as db:
            db.get(User, self.ids['admin']).role = 'VEHICLE_MANAGER'
            db.commit()
        self.assertEqual(self.request('GET', f'/mileage-reports/{job_id}/download').status_code, 403)

    def test_unauthorized_missing_and_expired_job(self):
        self.assertEqual(self.request('GET', '/mileage-reports/not-a-job').status_code, 404)
        job_id = 'b' * 32
        directory = report_jobs.root_path() / job_id
        directory.mkdir()
        report_jobs.write_state(directory / 'state.json', {'id': job_id, 'owner_id': self.ids['admin'], 'created_at': time.time() - report_jobs.TTL_SECONDS - 1})
        self.assertEqual(self.request('GET', '/mileage-reports/' + job_id).status_code, 410)

    def test_empty_report_is_valid_and_contains_no_fabricated_time(self):
        rows, filters = self.snapshot(driver_name='不存在的名字', photos='none')
        self.assertEqual(rows, [])
        path = Path(self.scratch.name) / 'empty.xlsx'
        stats = build_workbook(rows, filters.model_dump(mode='json'), self.uploads, path)
        self.assertEqual(stats['records'], 0)
        self.assertEqual(load_workbook(path)['个人里程汇总'].cell(7, 4).value, 0)


if __name__ == '__main__':
    unittest.main()
