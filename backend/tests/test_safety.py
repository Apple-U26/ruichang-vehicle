"""Run from backend: python -m unittest discover -s tests -v.

Uses an isolated in-memory database and real HTTP/authorization dependencies.
Never loads the deployment database or writes production data.
"""

import io
import os
import tempfile
import unittest
from datetime import date, datetime
from decimal import Decimal

os.environ['DATABASE_URL'] = 'sqlite://'
os.environ['JWT_SECRET'] = 'isolated-regression-tests-only-do-not-deploy'
UPLOADS = tempfile.TemporaryDirectory()
os.environ['UPLOAD_DIR'] = UPLOADS.name

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, event, func, select
from sqlalchemy.pool import StaticPool

from app import database
from app.models import (
    FuelRecord, MaintenanceRecord, MileageRecord, Project, Reimbursement,
    ReimbursementDetail, User, Vehicle, Welder,
)


def test_engine():
    engine = create_engine('sqlite://', poolclass=StaticPool, connect_args={'check_same_thread': False})

    @event.listens_for(engine, 'connect')
    def configure(connection, _):
        connection.execute('PRAGMA foreign_keys=ON')
        connection.create_function('date_format', 2, lambda value, fmt: str(value)[:7] if value else None)
    return engine


database.engine = test_engine()
database.SessionLocal.configure(bind=database.engine)
from app.main import app
from app.security import create_token, hash_password

PASSWORD_HASH = hash_password('TestPassword123!')


class SafetyTests(unittest.TestCase):
    def setUp(self):
        self.engine = test_engine()
        database.Base.metadata.create_all(self.engine)
        database.SessionLocal.configure(bind=self.engine)
        self.client = TestClient(app)
        self.client.__enter__()
        self.headers = {}
        self.ids = {}
        with database.SessionLocal() as db:
            for name, role in [('admin', 'ADMIN'), ('pm_a', 'PROJECT_MANAGER'), ('pm_b', 'PROJECT_MANAGER'), ('vm_a', 'VEHICLE_MANAGER'), ('finance', 'FINANCE'), ('driver', 'DRIVER'), ('unbound', 'DRIVER')]:
                user = User(username=name, real_name=name, password_hash=PASSWORD_HASH, role=role, enabled=True)
                db.add(user)
                db.flush()
                self.ids[name] = user.id
            pa = Project(name='Project A', manager_name='pm_a', manager_user_id=self.ids['pm_a'])
            pb = Project(name='Project B', manager_name='pm_b', manager_user_id=self.ids['pm_b'])
            db.add_all([pa, pb])
            db.flush()
            self.projects = {'A': pa.id, 'B': pb.id}
            self.vehicles = {}
            for index, name in enumerate('ABCD', 1):
                v = Vehicle(vehicle_code=f'CL-{index:06d}', plate_no='TEST-' + name, project_id=pa.id if name == 'A' else pb.id, vehicle_manager='vm_a' if name == 'A' else 'vm_b', status='ACTIVE')
                db.add(v)
                db.flush()
                self.vehicles[name] = v.id
            db.get(User, self.ids['driver']).vehicle_id = self.vehicles['A']
            db.commit()
            for user in db.scalars(select(User)):
                self.headers[user.username] = {'Authorization': 'Bearer ' + create_token(user)}

    def tearDown(self):
        self.client.__exit__(None, None, None)
        self.engine.dispose()

    def request(self, method, path, role='admin', **kwargs):
        return self.client.request(method, '/api' + path, headers=self.headers[role], **kwargs)

    def reimbursement(self, vehicle='A', details=None, role='admin'):
        return self.request('POST', '/reimbursements', role, json={
            'reimbursement_month': '2026-10', 'vehicle_id': self.vehicles[vehicle],
            'details': details or [{'expense_type': 'TOLL', 'expense_date': '2026-10-01', 'amount': 10, 'attachment_url': '/uploads/test.jpg'}],
        })

    def source_detail(self):
        with database.SessionLocal() as db:
            fuel = FuelRecord(vehicle_id=self.vehicles['A'], fuel_date=date(2026, 10, 1), total_amount=Decimal('100'), attachment_url='/uploads/test.jpg')
            db.add(fuel)
            db.commit()
            return {'expense_type': 'FUEL', 'expense_date': '2026-10-01', 'source_type': 'FUEL', 'source_id': fuel.id}

    def submitted(self, vehicle='A'):
        response = self.reimbursement(vehicle)
        self.assertEqual(response.status_code, 200, response.text)
        record_id = response.json()['id']
        self.assertEqual(self.request('POST', f'/reimbursements/{record_id}/submit').status_code, 200)
        return record_id

    def test_existing_login_and_health(self):
        self.assertEqual(self.request('GET', '/health').status_code, 200)
        response = self.client.post('/api/auth/login', json={'username': 'admin', 'password': 'TestPassword123!'})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(response.json()['access_token'])

    def test_consecutive_deletion_preserves_other_ids_codes_and_bindings(self):
        with database.SessionLocal() as db:
            driver = db.get(User, self.ids['driver'])
            driver.vehicle_id = self.vehicles['D']
            db.commit()
        for vehicle in 'BAC':
            response = self.request('DELETE', f'/vehicles/{self.vehicles[vehicle]}')
            self.assertEqual(response.status_code, 200, response.text)
        with database.SessionLocal() as db:
            remaining = db.scalars(select(Vehicle)).all()
            self.assertEqual([(v.id, v.plate_no, v.vehicle_code) for v in remaining], [(self.vehicles['D'], 'TEST-D', 'CL-000004')])
            self.assertEqual(db.get(User, self.ids['driver']).vehicle_id, self.vehicles['D'])

    def test_delete_bound_vehicle_unbinds_account(self):
        self.assertEqual(self.request('DELETE', f'/vehicles/{self.vehicles["A"]}').status_code, 200)
        with database.SessionLocal() as db:
            self.assertIsNone(db.get(User, self.ids['driver']).vehicle_id)

    def test_welder_delete_preserves_other_assets(self):
        with database.SessionLocal() as db:
            for name in 'ABC':
                db.add(Welder(welder_code='HJ-' + name, welder_no=name))
            db.commit()
        self.assertEqual(self.request('DELETE', '/welders/2').status_code, 200)
        self.assertEqual(self.request('DELETE', '/welders/1').status_code, 200)
        with database.SessionLocal() as db:
            self.assertEqual([(w.id, w.welder_code) for w in db.scalars(select(Welder))], [(3, 'HJ-C')])

    def test_duplicate_source_in_same_request_is_rejected_atomically(self):
        detail = self.source_detail()
        response = self.reimbursement(details=[detail, detail])
        self.assertEqual(response.status_code, 400, response.text)
        with database.SessionLocal() as db:
            self.assertEqual(db.scalar(select(func.count(Reimbursement.id))), 0)
            self.assertEqual(db.scalar(select(func.count(ReimbursementDetail.id))), 0)

    def test_source_cannot_be_reused_across_requests(self):
        detail = self.source_detail()
        self.assertEqual(self.reimbursement(details=[detail]).status_code, 200)
        self.assertEqual(self.reimbursement(details=[detail]).status_code, 400)

    def test_create_after_deleting_middle_reimbursement(self):
        rows = [self.reimbursement().json() for _ in range(3)]
        self.assertEqual(self.request('DELETE', f'/reimbursements/{rows[1]["id"]}').status_code, 200)
        created = self.reimbursement()
        self.assertEqual(created.status_code, 200, created.text)
        self.assertNotIn(created.json()['reimbursement_no'], [r['reimbursement_no'] for r in rows])

    def test_legacy_display_number_collision_is_retried(self):
        with database.SessionLocal() as db:
            db.add(Reimbursement(reimbursement_no='BX-202610-0002', reimbursement_month='2026-10', vehicle_id=self.vehicles['A'], applicant_id=self.ids['admin'], applicant_name='admin', total_amount=10))
            db.commit()
        response = self.reimbursement()
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['reimbursement_no'], 'BX-202610-0003')

    def test_own_project_can_approve_other_project_cannot(self):
        own = self.submitted('A')
        other = self.submitted('B')
        self.assertEqual(self.request('POST', f'/reimbursements/{other}/project-approve', 'pm_a', json={}).status_code, 403)
        self.assertEqual(self.request('POST', f'/reimbursements/{own}/project-approve', 'pm_a', json={}).status_code, 200)
        self.assertEqual(self.request('POST', f'/reimbursements/{own}/finance-approve', 'finance', json={}).status_code, 200)

    def test_mixed_batch_approval_rolls_back_when_scope_check_fails(self):
        own, other = self.submitted('A'), self.submitted('B')
        response = self.request('POST', '/reimbursements/batch-approve', 'pm_a', json={'ids': [own, other]})
        self.assertEqual(response.status_code, 403, response.text)
        with database.SessionLocal() as db:
            self.assertEqual(db.get(Reimbursement, own).status, 'SUBMITTED')
            self.assertEqual(db.get(Reimbursement, other).status, 'SUBMITTED')

    def test_outside_reimbursement_details_are_denied(self):
        record_id = self.reimbursement('B').json()['id']
        for role in ['pm_a', 'vm_a']:
            self.assertEqual(self.request('GET', f'/reimbursements/{record_id}', role).status_code, 403)

    def test_vehicle_update_requires_scope(self):
        for vehicle, expected in [('B', 403), ('A', 200)]:
            response = self.request('PUT', f'/vehicles/{self.vehicles[vehicle]}', 'pm_a', json={'plate_no': 'TEST-' + vehicle, 'project_id': self.projects[vehicle], 'vehicle_manager': 'vm_a'})
            self.assertEqual(response.status_code, expected, response.text)

    def test_unbound_driver_cannot_read_vehicle_or_export(self):
        self.assertEqual(self.request('GET', f'/vehicles/{self.vehicles["A"]}', 'unbound').status_code, 403)
        for path in ['/excel/export', '/excel/export/vehicles', '/excel/export/welders']:
            self.assertEqual(self.request('GET', path, 'driver').status_code, 403)

    def test_manager_export_only_contains_authorized_vehicles(self):
        for role in ['pm_a', 'vm_a']:
            response = self.request('GET', '/excel/export/vehicles', role)
            self.assertEqual(response.status_code, 200, response.text[:300] if response.status_code != 200 else '')
            workbook = load_workbook(io.BytesIO(response.content), read_only=True)
            plates = [row[1] for row in workbook['车辆信息'].iter_rows(min_row=2, values_only=True)]
            self.assertEqual(plates, ['TEST-A'])
        response = self.request('GET', '/excel/export/vehicles', 'admin')
        workbook = load_workbook(io.BytesIO(response.content), read_only=True)
        self.assertEqual(workbook['车辆信息'].max_row, 5)

    def test_full_export_aggregates_are_scoped(self):
        from app.excel_service import ExcelService
        from app.export_scope import export_scope
        with database.SessionLocal() as db:
            manager = db.get(User, self.ids['pm_a'])
            with export_scope(db, manager):
                self.assertEqual(ExcelService._get_dashboard_data(db)['overview']['总车辆数'], 1)
            self.assertEqual(ExcelService._get_dashboard_data(db)['overview']['总车辆数'], 4)
        response = self.request('GET', '/excel/export', 'pm_a')
        self.assertEqual(response.status_code, 200, response.text[:300] if response.status_code != 200 else '')

    def test_finance_cannot_register_trip_but_bound_driver_can(self):
        payload = {'vehicle_id': self.vehicles['A'], 'trip_date': '2026-10-24T10:00:00', 'out_mileage': 0, 'out_photo': '/uploads/out.jpg'}
        self.assertEqual(self.request('POST', '/mileages/out', 'finance', json=payload).status_code, 403)
        response = self.request('POST', '/mileages/out', 'driver', json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        response = self.request('PUT', f'/mileages/{response.json()["id"]}/close', 'driver', json={'in_mileage': 10, 'in_photo': '/uploads/in.jpg'})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['distance'], 10)

    def test_explicit_project_owner_id_prevents_name_based_override(self):
        with database.SessionLocal() as db:
            db.get(Project, self.projects['B']).manager_name = 'pm_a'
            db.commit()
        self.assertEqual(self.request('GET', f'/vehicles/{self.vehicles["B"]}', 'pm_a').status_code, 403)
        rows = self.request('GET', '/mileages', 'pm_a')
        self.assertEqual(rows.status_code, 200, rows.text)

    def test_non_admin_import_is_rejected(self):
        response = self.request('POST', '/excel/import', 'pm_a', files={'file': ('empty.xlsx', b'not-a-workbook')})
        self.assertEqual(response.status_code, 403, response.text)


if __name__ == '__main__':
    unittest.main()
